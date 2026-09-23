# -*- coding: utf-8 -*-
"""Search rows for the documents the site publishes -- the syllabus PDF and
the video slide decks (2026-09-23, Nico: "can we include search results for
things that are in the syllabus or other PDF (e.g., in specific slides)?").

Two kinds of row, and they behave differently on purpose:

  * SYLLABUS -- one row per PAGE. A hit links straight to that page, with
    the `#page=N` fragment every current browser's built-in PDF viewer
    honours.
  * SLIDES -- one row per SLIDE. A .pptx carries no way to address a slide
    from a URL, so the row NAMES the place ("see Module 3, Video 2, Slide
    14") and the link downloads the whole deck. Nico asked for exactly that
    rather than a jump that cannot work.

The decks are read from `C.VIDEO_SLIDES`, which is rescanned on every build,
so a deck added or replaced in "405 Slide Revisions 2026/Module N/Videos
Final" is indexed by the next `python _build_site.py` with nothing to
remember.

Text extraction is best-effort: a missing pypdf, a missing PDF or a deck
that will not open is REPORTED and skipped, never fatal. The website is the
deliverable; the search index is an enhancement to it.
"""

import os
import re

# Slide text only -- not the speaker notes. The notes are Nico's own script
# (asides, the answer to a discussion prompt, numbers deliberately kept off
# the slide), so indexing them would surface as "course content" things the
# class was never shown. Same reasoning as the podcast sourcing rule.
NOTES = False

SNIPPET = 90

# Every syllabus page opens with the running header and the page number, so
# a snippet taken from the top of the page read "MGMT 405 - Fall 2026 5
# encourage you to ..." -- the header first and the sentence mid-way. The
# full text still carries it; only the snippet is trimmed.
_HEADER = re.compile(r"^\s*MGMT\s*405\s*[·.\-|]*\s*Fall\s*2026\s*\d*\s*",
                     re.I)


def _tidy(s):
    return " ".join((s or "").split())


def _snippet(text, n=SNIPPET):
    text = _tidy(text)
    if len(text) <= n:
        return text
    cut = text[:n]
    sp = cut.rfind(" ")
    return (cut[:sp] if sp > n * 0.6 else cut).rstrip(" ,;:.-") + "…"


def syllabus_rows(pdf_path, url):
    """One row per page of the syllabus PDF."""
    if not os.path.exists(pdf_path):
        print("    search: no syllabus PDF at %s -- not indexed" % pdf_path)
        return []
    try:
        from pypdf import PdfReader
    except ImportError:
        print("    search: pypdf not installed -- the syllabus is not indexed")
        return []
    try:
        reader = PdfReader(pdf_path)
    except Exception as e:                                # noqa: BLE001
        print("    search: could not read the syllabus PDF (%s)" % e)
        return []

    rows = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = _tidy(page.extract_text())
        except Exception:                                 # noqa: BLE001
            text = ""
        if not text:
            continue
        rows.append({
            "href": "%s#page=%d" % (url, i),
            "kind": "Syllabus",
            "title": "Class Syllabus – page %d" % i,
            "sub": _snippet(_HEADER.sub("", text)),
            "text": text,
        })
    return rows


def _walk(shapes, out):
    """Collect text from a shape tree. These decks group heavily -- a curve
    with its label, a box with its text, a picture with its caption -- and
    groups nest, so this recurses rather than looking one level down."""
    for sh in shapes:
        try:
            if sh.shape_type == 6:            # MSO_SHAPE_TYPE.GROUP
                _walk(sh.shapes, out)
                continue
        except Exception:                                 # noqa: BLE001
            pass
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                out.extend(c.text for c in row.cells)


def _deck_title(path):
    """"Module 3 - Video 2 - The Production Function.pptx" -> the topic."""
    base = os.path.splitext(os.path.basename(path))[0]
    m = re.match(r"Module \d+ - Video \d+ - (.+)$", base)
    return m.group(1) if m else base


def slide_rows(video_slides, pub_name):
    """One row per slide of every video deck.

    `video_slides` is C.VIDEO_SLIDES -- {(module, video): absolute path} --
    and `pub_name` is C.slides_pub_name, which gives the space-free name the
    deck is served under.
    """
    try:
        from pptx import Presentation
    except ImportError:
        print("    search: python-pptx not installed -- slides not indexed")
        return []

    rows = []
    for (mod, vid), path in sorted(video_slides.items()):
        try:
            pres = Presentation(path)
        except Exception as e:                            # noqa: BLE001
            print("    search: could not read %s (%s)"
                  % (os.path.basename(path), e))
            continue
        topic = _deck_title(path)
        href = "slides/" + pub_name(path)
        for n, slide in enumerate(pres.slides, start=1):
            parts = []
            _walk(slide.shapes, parts)
            if NOTES and slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
            text = _tidy(" ".join(parts))
            if not text:
                continue
            rows.append({
                "href": href,
                "kind": "Slides",
                # the wording Nico asked for: name the place, do not pretend
                # the link can jump to it
                "title": "see Module %d, Video %d, Slide %d" % (mod, vid, n),
                "sub": "Download the full slide deck here",
                # the deck's topic joins the searchable text, so "production
                # function" finds its slides even where a given slide's own
                # text does not repeat the title
                "text": "%s %s" % (topic, text),
            })
    return rows
