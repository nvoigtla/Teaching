# -*- coding: utf-8 -*-
"""Display number -> "<top-bar tag> :: <action title>", for remapping the
number-keyed maps (animation plans, written notes, the splice map) after a
reordering.  The tag is part of the key because nine agenda slides share
the title "Outline of Module 6" and only the tag tells them apart."""
import json, sys
from pptx import Presentation
out = {}
prs = Presentation(sys.argv[1])
for i, sl in enumerate(prs.slides, 1):
    tag = title = ""
    for sh in sl.shapes:
        if not sh.has_text_frame or sh.left is None:
            continue
        x, y = sh.left / 914400.0, sh.top / 914400.0
        t = " ".join(sh.text_frame.text.split())
        if abs(x - 0.28) < 0.03 and abs(y - 0.55) < 0.04:
            title = t
        elif abs(x - 0.28) < 0.03 and y < 0.05:
            tag = t
    if not title:                  # dividers / title cards: all their text
        title = "|".join(" ".join(sh.text_frame.text.split())
                         for sh in sl.shapes if sh.has_text_frame)[:70]
    out[i] = "%s :: %s" % (tag, title)
json.dump(out, open(sys.argv[2], "w", encoding="utf-8"), ensure_ascii=False,
          indent=0)
print("%s: %d slides" % (sys.argv[1], len(out)))
